"""Builds the 'Freelancer & Small Business Finance Tracker' Excel template (US, USD)."""
import random
from datetime import date, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, PieChart, Reference
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.formatting.rule import CellIsRule

NAVY, TEAL, LIGHT, GREY = "1F2A44", "2A9D8F", "EAF6F4", "F3F4F6"
HDR = PatternFill("solid", fgColor=NAVY)
INPUT = PatternFill("solid", fgColor="FFF8E1")
CARD = PatternFill("solid", fgColor=LIGHT)
thin = Side(style="thin", color="D1D5DB")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
USD = '"$"#,##0.00'
MONTHS = ["Jan","Feb","Mar","Apr","May","Jun","Jul","Aug","Sep","Oct","Nov","Dec"]
INCOME_CATS = ["Client Services", "Product Sales", "Consulting", "Royalties", "Other Income"]
EXPENSE_CATS = ["Software & Subscriptions", "Advertising & Marketing", "Office Supplies",
                "Travel", "Meals (50%)", "Home Office", "Professional Services",
                "Equipment", "Insurance", "Education", "Fees & Processing", "Other"]
N = 500  # pre-formatted rows per ledger

wb = Workbook()

def header(ws, row, labels, widths=None):
    for i, l in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=l)
        c.font = Font(bold=True, color="FFFFFF"); c.fill = HDR
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    if widths:
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=row, column=i).column_letter].width = w

def title(ws, text, sub):
    ws["A1"] = text; ws["A1"].font = Font(size=20, bold=True, color=NAVY)
    ws["A2"] = sub; ws["A2"].font = Font(italic=True, color="6B7280")
    ws.sheet_view.showGridLines = False

# ---------- Settings ----------
st = wb.active; st.title = "Settings"
title(st, "Settings", "Yellow cells are yours to edit. Everything else calculates automatically.")
st["A4"], st["B4"] = "Business name", "My Business LLC"
st["A5"], st["B5"] = "Tax year", date.today().year
st["A6"], st["B6"] = "Estimated tax set-aside %", 0.25
st["B6"].number_format = "0%"
for r in (4,5,6):
    st[f"A{r}"].font = Font(bold=True); st[f"B{r}"].fill = INPUT; st[f"B{r}"].border = BOX
st["A8"] = "Income categories"; st["C8"] = "Expense categories"
st["A8"].font = st["C8"].font = Font(bold=True, color=NAVY)
for i, c in enumerate(INCOME_CATS): st.cell(row=9+i, column=1, value=c).fill = INPUT
for i, c in enumerate(EXPENSE_CATS): st.cell(row=9+i, column=3, value=c).fill = INPUT
st["A23"] = ("Tip: set-aside % is a planning aid only, not tax advice. "
             "Ask a CPA or use the IRS guidance for your situation.")
st["A23"].font = Font(italic=True, size=9, color="6B7280")
st.column_dimensions["A"].width = 30; st.column_dimensions["B"].width = 22
st.column_dimensions["C"].width = 30

# ---------- Ledgers ----------
def ledger(name, cats_col, cats_n, with_client):
    ws = wb.create_sheet(name)
    title(ws, name, "Add one row per transaction. Dates and categories are validated.")
    cols = ["Date", "Description", "Category", "Amount (USD)", "Month"]
    widths = [14, 38, 26, 16, 10]
    if with_client:
        cols += ["Client", "Invoice #", "Status"]; widths += [24, 14, 14]
    else:
        cols += ["Vendor", "Payment Method", "Tax Deductible?"]; widths += [24, 18, 16]
    header(ws, 4, cols, widths)
    ws.freeze_panes = "A5"
    dv = DataValidation(type="list", formula1=f"=Settings!${cats_col}$9:${cats_col}${8+cats_n}", allow_blank=True)
    ws.add_data_validation(dv)
    dd = DataValidation(type="date", operator="greaterThan", formula1="36526", allow_blank=True,
                        error="Enter a valid date", errorTitle="Invalid date")
    ws.add_data_validation(dd)
    extra = DataValidation(type="list", allow_blank=True,
        formula1='"Paid,Pending,Overdue"' if with_client else '"Yes,No,Partial"')
    ws.add_data_validation(extra)
    for r in range(5, 5+N):
        ws.cell(row=r, column=1).number_format = "mm/dd/yyyy"
        ws.cell(row=r, column=4).number_format = USD
        ws.cell(row=r, column=5, value=f'=IF(A{r}="","",MONTH(A{r}))')
        for c in range(1, len(cols)+1):
            cell = ws.cell(row=r, column=c); cell.border = BOX
            if c != 5: cell.fill = INPUT
        dv.add(ws.cell(row=r, column=3)); dd.add(ws.cell(row=r, column=1))
        extra.add(ws.cell(row=r, column=8))
    if with_client:
        red = PatternFill("solid", bgColor="FECACA")
        ws.conditional_formatting.add(f"H5:H{4+N}", CellIsRule(operator="equal", formula=['"Overdue"'], fill=red))
    return ws

inc = ledger("Income", "A", len(INCOME_CATS), True)
exp = ledger("Expenses", "C", len(EXPENSE_CATS), False)

# ---------- Sample data (clearly marked, easy to clear) ----------
random.seed(7)
y = date.today().year
clients = ["Acme Corp", "Brightside Studio", "Northwind LLC", "Evergreen Co", "Peak Digital"]
for i in range(36):
    r = 5 + i
    inc.cell(row=r, column=1, value=date(y, 1 + i % 12, random.randint(1, 27)))
    inc.cell(row=r, column=2, value=random.choice(["Website project", "Monthly retainer", "Consulting hours", "Template sale"]))
    inc.cell(row=r, column=3, value=random.choice(INCOME_CATS[:3]))
    inc.cell(row=r, column=4, value=round(random.uniform(350, 3200), 2))
    inc.cell(row=r, column=6, value=random.choice(clients))
    inc.cell(row=r, column=7, value=f"INV-{1001+i}")
    inc.cell(row=r, column=8, value=random.choice(["Paid"]*6 + ["Pending", "Overdue"]))
vendors = ["Adobe", "Google Ads", "Staples", "Delta", "Zoom", "QuickBooks", "Stripe", "Amazon"]
for i in range(60):
    r = 5 + i
    exp.cell(row=r, column=1, value=date(y, 1 + i % 12, random.randint(1, 27)))
    exp.cell(row=r, column=2, value="Sample expense")
    exp.cell(row=r, column=3, value=random.choice(EXPENSE_CATS))
    exp.cell(row=r, column=4, value=round(random.uniform(8, 650), 2))
    exp.cell(row=r, column=6, value=random.choice(vendors))
    exp.cell(row=r, column=7, value=random.choice(["Credit Card", "Debit Card", "ACH", "PayPal"]))
    exp.cell(row=r, column=8, value="Yes")

# ---------- Monthly Summary ----------
ms = wb.create_sheet("Monthly Summary")
title(ms, "Monthly Summary", "Auto-calculated from the Income and Expenses sheets.")
header(ms, 4, ["Month", "Income", "Expenses", "Net Profit", "Profit Margin", "Tax Set-Aside"],
       [12, 16, 16, 16, 15, 16])
rng_i = f"Income!$E$5:$E${4+N}"; amt_i = f"Income!$D$5:$D${4+N}"
rng_e = f"Expenses!$E$5:$E${4+N}"; amt_e = f"Expenses!$D$5:$D${4+N}"
for m in range(12):
    r = 5 + m
    ms.cell(row=r, column=1, value=MONTHS[m])
    ms.cell(row=r, column=2, value=f"=SUMIFS({amt_i},{rng_i},{m+1})")
    ms.cell(row=r, column=3, value=f"=SUMIFS({amt_e},{rng_e},{m+1})")
    ms.cell(row=r, column=4, value=f"=B{r}-C{r}")
    ms.cell(row=r, column=5, value=f'=IF(B{r}=0,0,D{r}/B{r})')
    ms.cell(row=r, column=6, value=f"=MAX(D{r},0)*Settings!$B$6")
    for c in range(1, 7):
        ms.cell(row=r, column=c).border = BOX
        if c in (2,3,4,6): ms.cell(row=r, column=c).number_format = USD
    ms.cell(row=r, column=5).number_format = "0.0%"
ms["A17"] = "TOTAL"
for c, L in zip(range(2, 7), "BCDEF"):
    ms.cell(row=17, column=c, value=f"=SUM({L}5:{L}16)" if L != "E" else "=IF(B17=0,0,D17/B17)")
    ms.cell(row=17, column=c).number_format = "0.0%" if L == "E" else USD
for c in range(1, 7):
    ms.cell(row=17, column=c).font = Font(bold=True); ms.cell(row=17, column=c).fill = CARD
    ms.cell(row=17, column=c).border = BOX

# ---------- Expense breakdown ----------
header(ms, 20, ["Expense Category", "Total", "% of Expenses"])
ms.column_dimensions["A"].width = 26
for i in range(len(EXPENSE_CATS)):
    r = 21 + i
    ms.cell(row=r, column=1, value=f"=Settings!C{9+i}")
    ms.cell(row=r, column=2, value=f'=SUMIFS({amt_e},Expenses!$C$5:$C${4+N},A{r})')
    ms.cell(row=r, column=3, value=f"=IF($C$17=0,0,B{r}/$C$17)")
    ms.cell(row=r, column=2).number_format = USD; ms.cell(row=r, column=3).number_format = "0.0%"
    for c in range(1, 4): ms.cell(row=r, column=c).border = BOX

# ---------- Dashboard ----------
db = wb.create_sheet("Dashboard", 0)
title(db, "=Settings!B4&\" — Finance Dashboard\"", "")
db["A1"] = '=Settings!B4&" — Finance Dashboard "&Settings!B5'
db["A2"] = "Replace the sample data in Income and Expenses with your own to update everything."
for c in "ABCDEFGHIJ": db.column_dimensions[c].width = 15
cards = [("Total Income", "='Monthly Summary'!B17", USD),
         ("Total Expenses", "='Monthly Summary'!C17", USD),
         ("Net Profit", "='Monthly Summary'!D17", USD),
         ("Profit Margin", "='Monthly Summary'!E17", "0.0%"),
         ("Suggested Tax Set-Aside", "='Monthly Summary'!F17", USD),
         ("Unpaid Invoices", f'=SUMIFS({amt_i},Income!$H$5:$H${4+N},"<>Paid")-SUMIFS({amt_i},Income!$H$5:$H${4+N},"")', USD)]
for i, (lab, f, fmt) in enumerate(cards):
    col = 1 + (i % 3) * 3; row = 4 + (i // 3) * 4
    db.merge_cells(start_row=row, start_column=col, end_row=row, end_column=col+1)
    db.merge_cells(start_row=row+1, start_column=col, end_row=row+2, end_column=col+1)
    a = db.cell(row=row, column=col, value=lab); a.font = Font(bold=True, color="FFFFFF"); a.fill = HDR
    a.alignment = Alignment(horizontal="center")
    db.cell(row=row, column=col+1).fill = HDR
    v = db.cell(row=row+1, column=col, value=f); v.number_format = fmt
    v.font = Font(size=18, bold=True, color=TEAL); v.alignment = Alignment(horizontal="center", vertical="center")
    for rr in (row+1, row+2):
        for cc in (col, col+1): db.cell(row=rr, column=cc).fill = CARD

bar = BarChart(); bar.type = "col"; bar.title = "Income vs Expenses by Month"
bar.add_data(Reference(ms, min_col=2, max_col=3, min_row=4, max_row=16), titles_from_data=True)
bar.set_categories(Reference(ms, min_col=1, min_row=5, max_row=16))
bar.height, bar.width = 8, 17
bar.series[0].graphicalProperties.solidFill = TEAL; bar.series[1].graphicalProperties.solidFill = "E76F51"
db.add_chart(bar, "A13")
pie = PieChart(); pie.title = "Expenses by Category"
pie.add_data(Reference(ms, min_col=2, min_row=20, max_row=20+len(EXPENSE_CATS)), titles_from_data=True)
pie.set_categories(Reference(ms, min_col=1, min_row=21, max_row=20+len(EXPENSE_CATS)))
pie.height, pie.width = 8, 12
db.add_chart(pie, "G13")
line = BarChart(); line.type = "col"; line.title = "Net Profit by Month"
line.add_data(Reference(ms, min_col=4, min_row=4, max_row=16), titles_from_data=True)
line.set_categories(Reference(ms, min_col=1, min_row=5, max_row=16))
line.legend = None; line.height, line.width = 7, 17
line.series[0].graphicalProperties.solidFill = NAVY
db.add_chart(line, "A30")

# ---------- Instructions ----------
ins = wb.create_sheet("Start Here", 0)
title(ins, "Start Here", "Freelancer & Small Business Finance Tracker — works in Excel and Google Sheets")
steps = ["1.  Open the Settings sheet: enter your business name, tax year and set-aside %.",
         "2.  Delete the sample rows in Income (A5:H40) and Expenses (A5:H64), keeping the formulas in the Month column.",
         "3.  Log each payment received in Income and each cost in Expenses (yellow cells only).",
         "4.  Pick categories from the dropdowns. Edit the category lists in Settings if you need to.",
         "5.  Watch the Dashboard update: profit, margin, unpaid invoices and tax set-aside.",
         "6.  Use Monthly Summary for bookkeeping hand-offs to your accountant at year end.",
         "",
         "Notes: Amounts are in USD. Meals are often only 50% deductible — confirm with your tax professional.",
         "This template is a bookkeeping aid and does not provide tax, legal or accounting advice."]
for i, s in enumerate(steps): 
    ins.cell(row=4+i, column=1, value=s).alignment = Alignment(wrap_text=True)
ins.column_dimensions["A"].width = 110

wb.move_sheet("Start Here", offset=-wb.index(wb["Start Here"]))
for ws in wb.worksheets:
    ws.sheet_properties.tabColor = TEAL if ws.title in ("Dashboard", "Start Here") else NAVY
wb.save("Freelancer_Finance_Tracker.xlsx")
print("saved")
